import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = docx.Document()

    # Set page margins (2 cm ~ 0.787 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style definitions
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(33, 33, 33)

    # Helper function for adding headings
    def add_custom_heading(text, level, color_rgb=(0, 51, 102)):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run = h.runs[0]
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(*color_rgb)
        return h

    def add_image_safely(img_path, caption, width=Inches(6.2)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=width)
            
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            crun = cp.add_run(caption)
            crun.font.italic = True
            crun.font.size = Pt(10.5)
            crun.font.color.rgb = RGBColor(90, 90, 90)
        else:
            p = doc.add_paragraph(f"[Image not found: {img_path}]")
            p.paragraph_format.space_after = Pt(6)

    # --- TITLE / HEADER ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    r_uni = title_p.add_run("TRƯỜNG ĐẠI HỌC VĂN LANG\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r_uni.font.size = Pt(13)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(0, 51, 102)

    star_p = doc.add_paragraph()
    star_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    star_p.paragraph_format.space_after = Pt(12)
    r_star = star_p.add_run("—" * 25)
    r_star.font.color.rgb = RGBColor(180, 180, 180)

    t_main = doc.add_paragraph()
    t_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_main.paragraph_format.space_after = Pt(14)
    r_main = t_main.add_run("BÁO CÁO BÀI TẬP THỰC HÀNH LAB 6\nTRIỂN KHAI VÀ QUẢN TRỊ ỨNG DỤNG TRÊN MICROSOFT AZURE")
    r_main.font.size = Pt(16)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0, 76, 153)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Sinh viên thực hiện:", "Võ Đình Giang"),
        ("Mã số sinh viên (MSSV):", "2374802010122"),
        ("Môn học:", "Các nền tảng phát triển phần mềm"),
        ("GitHub Repository:", "https://github.com/toolnayngonphet/azure-crm-api")
    ]
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(label)
        r_l.font.bold = True
        
        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(2)
        r_v = p_v.add_run(val)
        if "github" in val:
            r_v.font.color.rgb = RGBColor(0, 102, 204)
            r_v.font.underline = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- PHẦN 1 ---
    add_custom_heading("PHẦN 1: KẾT QUẢ THỰC HÀNH 4 MODULE MICROSOFT LEARN", level=1, color_rgb=(0, 76, 153))

    p_intro_1 = doc.add_paragraph(
        "Sinh viên đã hoàn thành đầy đủ cả 4 module đào tạo và kiểm tra kiến thức thực tế (Module Assessment) "
        "trên nền tảng Microsoft Learn với tài khoản sinh viên chính chủ (hiển thị logo avatar '2V' và thanh điểm XP/Level)."
    )
    p_intro_1.paragraph_format.space_after = Pt(8)

    # Module 1
    add_custom_heading("1. Module 1: Get started with Azure SQL Database for modern application development", level=2)
    p_m1 = doc.add_paragraph(
        "• Mục tiêu: Tìm hiểu kiến trúc Azure SQL Database, các mô hình triển khai (Single Database, Elastic Pool), "
        "tầng dịch vụ Hyperscale, cơ chế co giãn Serverless, tính năng tự phục hồi (Application Resilience) và tích hợp AI Copilot.\n"
        "• Kết quả: Hoàn thành 100% các đơn vị bài học và vượt qua bài đánh giá kiến thức Module Assessment."
    )
    p_m1.paragraph_format.space_after = Pt(4)
    img_m1 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293197301_0d5dfc53.png"
    add_image_safely(img_m1, "Hình 1.1: Chứng nhận hoàn thành Module 1 - Azure SQL Database trên Microsoft Learn")

    # Module 2
    add_custom_heading("2. Module 2: Host a web application with Azure App Service", level=2)
    p_m2 = doc.add_paragraph(
        "• Mục tiêu: Hiểu về nền tảng PaaS Azure App Service, cách đóng gói và triển khai ứng dụng web lên máy chủ Cloud, "
        "tìm hiểu cơ chế tự động co giãn tài nguyên (Autoscale) và các nguồn triển khai tự động (Continuous Deployment).\n"
        "• Kết quả: Đạt chứng nhận hoàn thành Module Assessment."
    )
    p_m2.paragraph_format.space_after = Pt(4)
    img_m2 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293202023_e18ee3bb.png"
    add_image_safely(img_m2, "Hình 1.2: Chứng nhận hoàn thành Module 2 - Host a web app with Azure App Service")

    # Module 3
    add_custom_heading("3. Module 3: Monitor applications with Application Insights", level=2)
    p_m3 = doc.add_paragraph(
        "• Mục tiêu: Nắm vững giải pháp giám sát hiệu năng APM (Application Performance Monitoring) với Application Insights, "
        "cấu hình thu thập log, metrics (Standard metrics, Platform metrics), tỉ lệ lấy mẫu (Sampling) và theo dõi Live Metrics.\n"
        "• Kết quả: Vượt qua bài kiểm tra kiến thức và nhận huy hiệu Module Assessment."
    )
    p_m3.paragraph_format.space_after = Pt(4)
    img_m3 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293206984_00b7214b.png"
    add_image_safely(img_m3, "Hình 1.3: Chứng nhận hoàn thành Module 3 - Configure monitoring for applications")

    # Module 4
    add_custom_heading("4. Module 4: Build and deploy applications to Azure by using GitHub Actions", level=2)
    p_m4 = doc.add_paragraph(
        "• Mục tiêu: Thiết lập pipeline CI/CD tự động hóa, cấu hình trigger sự kiện, quản lý bí mật an toàn với GitHub Secrets "
        "và triển khai ứng dụng sang môi trường Azure Web App.\n"
        "• Kết quả: Hoàn thành bài đánh giá kiến thức đạt cấp độ Level 3 trên Microsoft Learn."
    )
    p_m4.paragraph_format.space_after = Pt(4)
    img_m4 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293211428_4f6a4f89.png"
    add_image_safely(img_m4, "Hình 1.4: Chứng nhận hoàn thành Module 4 - Build & deploy applications using GitHub Actions")

    # --- PHẦN 2 ---
    add_custom_heading("PHẦN 2: THỰC HÀNH MỞ RỘNG - XÂY DỰNG VÀ TRIỂN KHAI HỆ THỐNG CRM API", level=1, color_rgb=(0, 76, 153))

    p_p2_desc = doc.add_paragraph(
        "Áp dụng các kiến thức đã học từ 4 module, sinh viên đã xây dựng hoàn chỉnh một ứng dụng Backend REST API "
        "quản lý khách hàng (CRM Customer Management) với đầy đủ quy trình DevOps hiện đại."
    )
    p_p2_desc.paragraph_format.space_after = Pt(8)

    # 2.1 Kiến trúc
    add_custom_heading("1. Kiến trúc hệ thống và Công nghệ sử dụng", level=2)
    tech_table = doc.add_table(rows=6, cols=2)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_items = [
        ("Nền tảng Backend", "Node.js (v24 LTS) kết hợp Express Framework & TypeScript"),
        ("Tài liệu & Test API", "Swagger UI (OpenAPI Specification 3.0) tích hợp tại endpoint /api-docs"),
        ("Tự động hóa CI/CD", "GitHub Actions với workflow .github/workflows/azure-deploy.yml"),
        ("Hạ tầng Cloud Hosting", "Azure App Service (Linux Web App, Node runtime)"),
        ("Cơ chế Co giãn (Scale)", "Quy tắc Autoscale: Scale Out khi CPU > 70%, Scale In khi CPU < 30%"),
        ("Giám sát Telemetry", "Azure Application Insights SDK (theo dõi requests, uptime, latency)")
    ]
    for idx, (comp, desc) in enumerate(tech_items):
        r = tech_table.rows[idx]
        c1, c2 = r.cells[0], r.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        c1.paragraphs[0].add_run(comp).font.bold = True
        c2.paragraphs[0].add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2.2 GitHub Repository
    add_custom_heading("2. Quản lý mã nguồn và Pipeline CI/CD trên GitHub", level=2)
    p_git = doc.add_paragraph(
        "Toàn bộ source code của dự án đã được khởi tạo Git và đẩy lên GitHub Repository công khai tại:\n"
        "🔗 https://github.com/toolnayngonphet/azure-crm-api\n\n"
        "Quy trình CI/CD được thiết lập trong file `.github/workflows/azure-deploy.yml`: Mỗi khi lập trình viên "
        "thực hiện lệnh git push lên nhánh `main`, GitHub Actions sẽ tự động kích hoạt tiến trình:\n"
        "  1. Checkout source code từ GitHub.\n"
        "  2. Thiết lập môi trường Node.js và cài đặt dependencies (`npm ci`).\n"
        "  3. Biên dịch TypeScript sang JavaScript (`npm run build`).\n"
        "  4. Đóng gói sản phẩm và deploy sang Azure Web App thông qua Publish Profile."
    )
    p_git.paragraph_format.space_after = Pt(4)
    img_git = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293856240_221df157.png"
    add_image_safely(img_git, "Hình 2.1: Repository dự án azure-crm-api được khởi tạo trên GitHub")

    # 2.3 Kết quả kiểm thử API CRM
    add_custom_heading("3. Kết quả Kiểm thử Chức năng API CRM (CRUD Khách hàng)", level=2)
    
    p_api_desc = doc.add_paragraph(
        "Hệ thống cung cấp đầy đủ các endpoint RESTful để quản lý thông tin khách hàng, cho phép các ứng dụng "
        "Frontend Web hoặc Mobile App dễ dàng kết nối và thao tác dữ liệu:"
    )
    p_api_desc.paragraph_format.space_after = Pt(4)

    # Image 1: Swagger overview
    img_swg = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293656029_f47db995.png"
    add_image_safely(img_swg, "Hình 2.2: Giao diện trực quan tài liệu Swagger UI của hệ thống Azure CRM API")

    # Image 2: GET customers
    p_get = doc.add_paragraph(
        "• Chức năng Xem danh sách khách hàng (`GET /api/customers`): Hỗ trợ tìm kiếm theo tên, email, "
        "công ty và lọc theo trạng thái (Lead, Contacted, Customer, Inactive). Kết quả trả về danh sách khách hàng "
        "cùng mã trạng thái HTTP 200 OK."
    )
    p_get.paragraph_format.space_after = Pt(4)
    img_get = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293692808_95f4f5c3.png"
    add_image_safely(img_get, "Hình 2.3: Kiểm thử chức năng GET /api/customers trả về danh sách khách hàng thành công (HTTP 200)")

    # Image 3: POST customer
    p_post = doc.add_paragraph(
        "• Chức năng Thêm khách hàng mới (`POST /api/customers`): Cho phép ứng dụng gửi payload JSON chứa "
        "thông tin khách hàng. Hệ thống tự động sinh mã định danh (CUST-xxx) và trả về thông tin khách hàng mới kèm mã HTTP 201 Created."
    )
    p_post.paragraph_format.space_after = Pt(4)
    img_post = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293718693_2bb2487f.png"
    add_image_safely(img_post, "Hình 2.4: Kiểm thử chức năng POST /api/customers tạo khách hàng mới thành công (HTTP 201)")

    # Image 4: Health check
    p_health = doc.add_paragraph(
        "• Endpoint Kiểm tra sức khỏe hệ thống (`GET /api/health`): Cung cấp trạng thái UP, thời gian hoạt động "
        "(uptime) và môi trường. Endpoint này phục vụ trực tiếp cho Azure Health Probe và cơ chế Autoscale của App Service."
    )
    p_health.paragraph_format.space_after = Pt(4)
    img_health = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293732866_08307b4b.png"
    add_image_safely(img_health, "Hình 2.5: Endpoint Health Check (/api/health) phản hồi trạng thái hoạt động của máy chủ")

    # 2.4 Autoscale & App Insights
    add_custom_heading("4. Cấu hình Autoscale và Giám sát Hiệu năng (Monitoring)", level=2)
    p_scale = doc.add_paragraph(
        "• Cấu hình Autoscale (Tự động co giãn):\n"
        "  - Mục đích: Đảm bảo ứng dụng đáp ứng tốt lượng truy cập đột biến trong khi vẫn tối ưu chi phí hạ tầng.\n"
        "  - Quy tắc Scale Out (Mở rộng): Khi mức sử dụng CPU trung bình vượt ngưỡng 70% trong 5 phút, "
        "hệ thống tự động kích hoạt thêm 1 instance máy chủ mới (tối đa 3 instances).\n"
        "  - Quy tắc Scale In (Thu hẹp): Khi CPU giảm xuống dưới 30% trong 10 phút, hệ thống tự động giảm bớt "
        "instance dư thừa về mức tối thiểu (1 instance).\n\n"
        "• Giám sát với Azure Application Insights:\n"
        "  - Ứng dụng đã được tích hợp sẵn SDK `applicationinsights`.\n"
        "  - Toàn bộ thông số về thời gian phản hồi (duration), lưu lượng request, tỉ lệ lỗi và logs console "
        "đều được tự động ghi nhận và đồng bộ trực tiếp lên Dashboard Application Insights trên Azure."
    )
    p_scale.paragraph_format.space_after = Pt(8)

    # --- PHẦN 3: KẾT LUẬN ---
    add_custom_heading("PHẦN 3: KẾT LUẬN VÀ BÀI HỌC KINH NGHIỆM", level=1, color_rgb=(0, 76, 153))
    p_concl = doc.add_paragraph(
        "Qua bài thực hành Lab 6, sinh viên đã đạt được các mục tiêu trọng tâm:\n"
        "1. Nắm vững cách thức hoạt động và triển khai các dịch vụ PaaS hàng đầu của Microsoft Azure: Azure App Service, "
        "Azure SQL Database, Application Insights.\n"
        "2. Xây dựng và tích hợp thành công chu trình DevOps khép kín: Viết mã nguồn Backend -> Quản lý phiên bản trên GitHub -> "
        "Tự động hóa Build & Deploy bằng GitHub Actions -> Giám sát hiệu năng liên tục bằng Application Insights.\n"
        "3. Hiểu rõ tư duy thiết kế hệ thống Cloud-native với khả năng mở rộng linh hoạt (Scalability & Resilience), "
        "sẵn sàng cung cấp các dịch vụ API chất lượng cao phục vụ ứng dụng Web và Mobile trong thực tế doanh nghiệp."
    )
    p_concl.paragraph_format.space_after = Pt(12)

    # Save document
    output_path = r"d:\HK261\Các nền tảng phát triển phần mềm\Lab6_ Azure App\Bao_Cao_Lab6_Azure_VoDinhGiang.docx"
    doc.save(output_path)
    print("Report generated successfully!")

if __name__ == "__main__":
    create_report()
