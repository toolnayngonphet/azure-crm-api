import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_border(cell, **kwargs):
    """
    Set cell borders: top, bottom, left, right
    kwargs can be top={"sz": 4, "val": "single", "color": "CCCCCC"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def create_report():
    doc = docx.Document()

    # Standard Vietnamese academic page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.8)     # ~2.0 cm
    section.bottom_margin = Inches(0.8)  # ~2.0 cm
    section.left_margin = Inches(1.0)    # ~2.5 cm
    section.right_margin = Inches(0.8)   # ~2.0 cm

    # Base style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # Helper: Normal Body Paragraph
    def add_body_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # Helper: Headings in BLACK, Standard Style
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # Helper: Add image safely with clean caption
    def add_image_figure(img_path, caption_text, width=Inches(5.8)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run()
            run.add_picture(img_path, width=width)

            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            cp.paragraph_format.first_line_indent = Inches(0)
            crun = cp.add_run(caption_text)
            crun.font.name = 'Times New Roman'
            crun.font.size = Pt(11)
            crun.font.italic = True
            crun.font.color.rgb = RGBColor(50, 50, 50)
        else:
            p = doc.add_paragraph(f"[Ảnh: {caption_text}]")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # =========================================================================
    # TRANG BÌA (COVER PAGE)
    # =========================================================================
    p_cov_top = doc.add_paragraph()
    p_cov_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_top.paragraph_format.space_after = Pt(2)
    p_cov_top.paragraph_format.first_line_indent = Inches(0)
    r1 = p_cov_top.add_run("TRƯỜNG ĐẠI HỌC VĂN LANG\n")
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_cov_top.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r2.font.size = Pt(13)
    r2.font.bold = True
    r3 = p_cov_top.add_run("BỘ MÔN CÔNG NGHỆ PHẦN MỀM")
    r3.font.size = Pt(12)

    p_cov_line = doc.add_paragraph()
    p_cov_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_line.paragraph_format.space_before = Pt(4)
    p_cov_line.paragraph_format.space_after = Pt(45)
    p_cov_line.paragraph_format.first_line_indent = Inches(0)
    p_cov_line.add_run("———————————————").font.color.rgb = RGBColor(100, 100, 100)

    p_cov_title = doc.add_paragraph()
    p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_title.paragraph_format.space_after = Pt(12)
    p_cov_title.paragraph_format.first_line_indent = Inches(0)
    r_bt = p_cov_title.add_run("BÁO CÁO THỰC HÀNH LAB 6\n")
    r_bt.font.size = Pt(18)
    r_bt.font.bold = True

    p_cov_sub = doc.add_paragraph()
    p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_sub.paragraph_format.space_after = Pt(50)
    p_cov_sub.paragraph_format.first_line_indent = Inches(0)
    r_bs = p_cov_sub.add_run("TRIỂN KHAI VÀ QUẢN TRỊ ỨNG DỤNG\nTRÊN NỀN TẢNG CLOUD MICROSOFT AZURE")
    r_bs.font.size = Pt(15)
    r_bs.font.bold = True

    # Khung thông tin sinh viên trên trang bìa
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(1.2)
    p_info.paragraph_format.line_spacing = 1.4
    p_info.paragraph_format.space_after = Pt(70)
    p_info.paragraph_format.first_line_indent = Inches(0)

    runs_info = [
        ("Môn học:\t\t", True), ("Các nền tảng phát triển phần mềm\n", False),
        ("Sinh viên thực hiện:\t", True), ("Võ Đình Giang\n", False),
        ("Mã số sinh viên:\t", True), ("2374802010122\n", False),
        ("Lớp học phần:\t", True), ("71ITDS30103\n", False),
        ("GitHub Repository:\t", True), ("https://github.com/toolnayngonphet/azure-crm-api", False)
    ]
    for text, is_bold in runs_info:
        r = p_info.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = is_bold

    p_cov_bot = doc.add_paragraph()
    p_cov_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_bot.paragraph_format.first_line_indent = Inches(0)
    r_city = p_cov_bot.add_run("TP. HỒ CHÍ MINH, THÁNG 10/2026")
    r_city.font.size = Pt(12)
    r_city.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # NỘI DUNG CHÍNH BÁO CÁO
    # =========================================================================

    add_h1("I. GIỚI THIỆU TỔNG QUAN BÀI THỰC HÀNH")
    add_body_p(
        "Mục tiêu của bài thực hành Lab 6 là giúp sinh viên nắm vững quy trình thiết kế, triển khai, "
        "tự động hóa và giám sát một ứng dụng web trên nền tảng điện toán đám mây Microsoft Azure. "
        "Bài thực hành bao gồm hai nội dung chính:"
    )
    add_body_p(
        "1. Hoàn thành 4 module chuyên sâu trên Microsoft Learn về Azure SQL Database, Azure App Service, "
        "Application Insights và tự động hóa triển khai bằng GitHub Actions."
    )
    add_body_p(
        "2. Áp dụng xây dựng hệ thống Backend API quản lý khách hàng (CRM), tích hợp tài liệu trực quan Swagger UI, "
        "thiết lập quy trình CI/CD qua GitHub Actions, cấu hình cơ chế tự động co giãn (Autoscale) và giám sát hiệu năng hệ thống."
    )

    # -------------------------------------------------------------------------
    add_h1("II. KẾT QUẢ THỰC HÀNH 4 MODULE TRÊN MICROSOFT LEARN")
    add_body_p(
        "Em đã hoàn thành học tập và vượt qua các bài đánh giá kiến thức (Module Assessment) của cả 4 module đào tạo "
        "trên hệ thống Microsoft Learn bằng tài khoản sinh viên chính chủ."
    )

    # 2.1 Module 1
    add_h2("1. Module 1: Get started with Azure SQL Database for modern application development")
    add_body_p(
        "Trong module này, em đã tìm hiểu về các tùy chọn triển khai cơ sở dữ liệu trên Azure (Single Database và Elastic Pool), "
        "kiến trúc tầng lưu trữ phân tán Hyperscale, mô hình tính toán Serverless và cơ chế đảm bảo tính sẵn sàng "
        "(High Availability & Disaster Recovery) với tính năng Active Geo-Replication và Auto-failover groups."
    )
    img_m1 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293197301_0d5dfc53.png"
    add_image_figure(img_m1, "Hình 1: Kết quả hoàn thành Module 1 - Azure SQL Database")

    # 2.2 Module 2
    add_h2("2. Module 2: Host a web application with Azure App Service")
    add_body_p(
        "Module giúp em hiểu rõ mô hình điện toán PaaS (Platform as a Service) với Azure App Service. "
        "Dịch vụ này cho phép lập trình viên tập trung vào phát triển mã nguồn ứng dụng web mà không cần bận tâm "
        "về việc quản trị hệ điều hành hay phần cứng bên dưới. Đồng thời, module cung cấp kiến thức về cách App Service "
        "tự động co giãn tài nguyên (Autoscale) và hỗ trợ tích hợp triển khai liên tục từ GitHub và Azure Repos."
    )
    img_m2 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293202023_e18ee3bb.png"
    add_image_figure(img_m2, "Hình 2: Kết quả hoàn thành Module 2 - Host a web app with Azure App Service")

    # 2.3 Module 3
    add_h2("3. Module 3: Monitor applications with Application Insights")
    add_body_p(
        "Module này cung cấp giải pháp giám sát hiệu năng ứng dụng (APM). Qua bài học, em nắm được cách thức "
        "Application Insights thu thập dữ liệu đo từ xa (Telemetry), phân biệt giữa Resource Logs và Activity Logs, "
        "hiểu ưu điểm về hiệu năng của Standard metrics trong thời gian truy vấn và cách áp dụng cơ chế Fixed-rate sampling "
        "để tối ưu lưu lượng truyền tải giữa trình duyệt người dùng và máy chủ."
    )
    img_m3 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293206984_00b7214b.png"
    add_image_figure(img_m3, "Hình 3: Kết quả hoàn thành Module 3 - Configure monitoring for applications")

    # 2.4 Module 4
    add_h2("4. Module 4: Build and deploy applications to Azure by using GitHub Actions")
    add_body_p(
        "Nội dung module hướng dẫn tự động hóa quy trình phát hành phần mềm bằng GitHub Actions. Em đã thực hành "
        "cấu hình workflow triển khai, quản lý khóa bảo mật an toàn thông qua GitHub Secrets (không để lộ thông tin xác thực "
        "dưới dạng văn bản thuần trong repository) và kiểm soát các điều kiện trigger tự động khi có commit mã nguồn mới."
    )
    img_m4 = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293211428_4f6a4f89.png"
    add_image_figure(img_m4, "Hình 4: Kết quả hoàn thành Module 4 - Build and deploy applications using GitHub Actions")

    # -------------------------------------------------------------------------
    add_h1("III. XÂY DỰNG VÀ TRIỂN KHAI ỨNG DỤNG CRM (CRUD KHÁCH HÀNG)")

    add_h2("1. Mô tả bài toán và Kiến trúc kỹ thuật")
    add_body_p(
        "Nhằm áp dụng kiến thức vào thực tế phát triển phần mềm doanh nghiệp, em xây dựng một hệ thống RESTful API "
        "phục vụ phân hệ quản lý khách hàng (CRM Customer Management). Ứng dụng cung cấp các dịch vụ để ứng dụng di động "
        "(Mobile App) hoặc ứng dụng Web Frontend kết nối và thao tác dữ liệu khách hàng."
    )
    add_body_p("Các công nghệ chính được lựa chọn triển khai bao gồm:")
    add_body_p("• Ngôn ngữ & Framework: Node.js, Express Framework kết hợp TypeScript giúp mã nguồn chặt chẽ và an toàn về kiểu dữ liệu.")
    add_body_p("• Tài liệu hóa API: Swagger UI tích hợp chuẩn OpenAPI 3.0 trực quan ngay tại đường dẫn /api-docs.")
    add_body_p("• Giám sát hệ thống: Tích hợp thư viện SDK applicationinsights để ghi nhận log, metrics và thời gian xử lý request.")
    add_body_p("• Tự động hóa CI/CD: Pipeline GitHub Actions (.github/workflows/azure-deploy.yml).")

    # 3.2 GitHub Repo
    add_h2("2. Quản lý mã nguồn và Tự động hóa CI/CD với GitHub Actions")
    add_body_p(
        "Toàn bộ source code dự án được khởi tạo và đẩy lên GitHub Repository công khai:"
    )
    add_body_p("🔗 Link Repository: https://github.com/toolnayngonphet/azure-crm-api")
    add_body_p(
        "File workflow `.github/workflows/azure-deploy.yml` được cấu hình để mỗi khi nhánh `main` nhận commit mới, "
        "hệ thống tự động thực hiện các bước: Cài đặt thư viện dependencies (`npm ci`) -> Biên dịch TypeScript (`npm run build`) -> "
        "Đóng gói artifact -> Triển khai lên Azure Web App thông qua bí mật AZURE_WEBAPP_PUBLISH_PROFILE được lưu trong GitHub Secrets."
    )
    img_git = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293856240_221df157.png"
    add_image_figure(img_git, "Hình 5: Khởi tạo và đồng bộ mã nguồn dự án lên GitHub Repository")

    # 3.3 API Details & Testing
    add_h2("3. Kết quả kiểm thử các Endpoint API CRM (CRUD Khách hàng)")
    add_body_p(
        "Hệ thống cung cấp đầy đủ các endpoint xử lý dữ liệu khách hàng với định dạng JSON tiêu chuẩn. "
        "Dưới đây là kết quả kiểm thử thực tế từng chức năng:"
    )

    # Table of endpoints
    table_api = doc.add_table(rows=7, cols=3)
    table_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Phương thức", "Đường dẫn (Endpoint)", "Chức năng"]
    api_rows = [
        ("GET", "/api/health", "Kiểm tra tình trạng hoạt động (Health probe cho Autoscale)"),
        ("GET", "/api/customers", "Lấy danh sách tất cả khách hàng (hỗ trợ tìm kiếm, lọc)"),
        ("GET", "/api/customers/{id}", "Xem thông tin chi tiết một khách hàng theo ID"),
        ("POST", "/api/customers", "Tạo mới một khách hàng trong hệ thống CRM"),
        ("PUT", "/api/customers/{id}", "Cập nhật thông tin khách hàng"),
        ("DELETE", "/api/customers/{id}", "Xóa khách hàng khỏi hệ thống")
    ]
    # Header row formatting
    hdr_cells = table_api.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(11)
        set_cell_shading(hdr_cells[i], "F0F0F0")
        set_cell_border(hdr_cells[i], bottom={"sz": 6, "val": "single", "color": "999999"})

    for row_idx, (m, ep, d) in enumerate(api_rows, start=1):
        row = table_api.rows[row_idx]
        row.cells[0].text = m
        row.cells[1].text = ep
        row.cells[2].text = d
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(11)
            set_cell_border(cell, bottom={"sz": 4, "val": "single", "color": "E0E0E0"})

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Image: Swagger
    add_h3("a. Giao diện Swagger UI trực quan")
    add_body_p(
        "Tài liệu API được sinh tự động và hiển thị tại `/api-docs`, cho phép lập trình viên frontend hoặc mobile "
        "dễ dàng tra cứu thông số kỹ thuật và gửi request kiểm thử trực tiếp trên trình duyệt."
    )
    img_swg = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293656029_f47db995.png"
    add_image_figure(img_swg, "Hình 6: Giao diện Swagger UI của hệ thống Azure CRM API")

    # Image: GET
    add_h3("b. Kiểm thử chức năng Xem danh sách khách hàng (GET /api/customers)")
    add_body_p(
        "Khi gọi endpoint `GET /api/customers`, hệ thống phản hồi mã trạng thái 200 OK kèm theo danh sách mảng JSON "
        "chứa các thông tin: ID, họ tên, email, số điện thoại, công ty, trạng thái (Lead, Customer, Contacted) và ngày tạo."
    )
    img_get = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293692808_95f4f5c3.png"
    add_image_figure(img_get, "Hình 7: Kết quả kiểm thử endpoint GET /api/customers thành công (HTTP 200)")

    # Image: POST
    add_h3("c. Kiểm thử chức năng Thêm mới khách hàng (POST /api/customers)")
    add_body_p(
        "Gửi request POST kèm thông tin khách hàng mới 'Đặng Minh Khang'. Hệ thống xử lý kiểm tra tính hợp lệ dữ liệu, "
        "tự động sinh mã CUST-004 và trả về thông báo 'Customer created successfully' cùng mã phản hồi HTTP 201 Created."
    )
    img_post = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293718693_2bb2487f.png"
    add_image_figure(img_post, "Hình 8: Kết quả kiểm thử endpoint POST /api/customers tạo mới khách hàng (HTTP 201)")

    # Image: Health
    add_h3("d. Kiểm thử endpoint Health Check (GET /api/health)")
    add_body_p(
        "Endpoint `/api/health` trả về trạng thái UP, thời gian hoạt động liên tục (uptime) và môi trường. "
        "Dữ liệu này cung cấp cho bộ giám sát kiểm tra tính sẵn sàng của phiên bản ứng dụng."
    )
    img_health = "C:/Users/Giang/.gemini/antigravity/brain/4eb5f8a9-0a6d-472b-9d1f-460cff0e16f3/.user_uploaded/media_1792293732866_08307b4b.png"
    add_image_figure(img_health, "Hình 9: Kết quả kiểm thử endpoint kiểm tra sức khỏe hệ thống (/api/health)")

    # 3.4 Autoscale & App Insights
    add_h2("4. Cấu hình Autoscale và Giám sát Hiệu năng trên Azure")
    add_body_p(
        "• Cơ chế tự động co giãn (Autoscale): Trong môi trường Azure App Service Plan, hệ thống được thiết lập quy tắc "
        "co giãn tự động dựa trên tải CPU thực tế nhằm vừa đảm bảo hiệu năng phục vụ người dùng, vừa tối ưu chi phí hạ tầng:"
    )
    add_body_p(
        "  + Quy tắc Scale Out (Mở rộng): Khi chỉ số CPU trung bình vượt quá 70% trong khoảng thời gian 5 phút liên tục, "
        "hệ thống tự động bổ sung thêm 1 instance máy chủ mới (giới hạn tối đa 3 instances)."
    )
    add_body_p(
        "  + Quy tắc Scale In (Thu hẹp): Khi mức tải CPU giảm xuống dưới 30% trong 10 phút, hệ thống tự động giải phóng "
        "instance dư thừa để quay về số lượng tối thiểu (1 instance)."
    )
    add_body_p(
        "• Giám sát với Application Insights: Ứng dụng được cấu hình chuỗi kết nối qua biến môi trường "
        "`APPLICATIONINSIGHTS_CONNECTION_STRING`. Nhờ đó, mọi truy cập API, thời gian phản hồi từng request và các ngoại lệ "
        "đều được tự động đồng bộ lên dịch vụ giám sát của Azure để theo dõi Live Metrics theo thời gian thực."
    )

    # -------------------------------------------------------------------------
    add_h1("IV. KẾT LUẬN")
    add_body_p(
        "Sau khi hoàn thành bài thực hành Lab 6, em đã nắm vững các kiến thức cốt lõi và kỹ năng thực tiễn sau:"
    )
    add_body_p(
        "1. Hiểu rõ cách thức tổ chức, cấu hình và vận hành các dịch vụ đám mây PaaS của Azure bao gồm App Service, "
        "Azure SQL Database và Application Insights."
    )
    add_body_p(
        "2. Áp dụng thành công quy trình DevOps: quản lý mã nguồn trên GitHub, tự động hóa build và triển khai qua GitHub Actions "
        "giúp rút ngắn thời gian phát hành phần mềm."
    )
    add_body_p(
        "3. Xây dựng hoàn chỉnh hệ thống REST API phục vụ chức năng CRM, hiểu rõ cơ chế co giãn tài nguyên tự động (Autoscale) "
        "giúp ứng dụng duy trì tính sẵn sàng cao khi có sự gia tăng đột biến về lượng truy cập."
    )

    # Save
    out_file = r"d:\HK261\Các nền tảng phát triển phần mềm\Lab6_ Azure App\Bao_Cao_Lab6_Azure_VoDinhGiang.docx"
    doc.save(out_file)
    print("Clean academic docx generated successfully!")

if __name__ == "__main__":
    create_report()
